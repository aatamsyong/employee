import os
import re
from typing import List, Dict, Any

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    Document = None

try:
    from pptx import Presentation
    from pptx.util import Inches as PptxInches, Pt as PptxPt
    from pptx.dml.color import RGBColor as PptxRGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    Presentation = None

def generate_docx(title: str, subtitle: str, sections: List[Dict[str, Any]], output_path: str) -> str:
    """Generates a professional DOCX proposal or SOP document."""
    if Document is None:
        raise ImportError("python-docx is not installed.")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.name = "Calibri"
    title_run.font.color.rgb = RGBColor(16, 44, 87)

    if subtitle:
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(subtitle)
        sub_run.font.size = Pt(13)
        sub_run.font.italic = True
        sub_run.font.name = "Calibri"
        sub_run.font.color.rgb = RGBColor(80, 80, 80)

    # Decorative separator
    divider = doc.add_paragraph()
    divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    div_run = divider.add_run("―" * 40)
    div_run.font.color.rgb = RGBColor(200, 200, 200)

    # Add Sections
    for sec in sections:
        sec_title = sec.get("title", "")
        sec_content = sec.get("content", "")
        sec_bullets = sec.get("bullets", [])
        sec_table = sec.get("table", None)

        if sec_title:
            h = doc.add_heading(sec_title, level=1)
            h_run = h.runs[0] if h.runs else h.add_run(sec_title)
            h_run.font.size = Pt(14)
            h_run.font.name = "Calibri"
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(27, 79, 114)

        if sec_content:
            paragraphs = sec_content.split("\n\n")
            for para_text in paragraphs:
                if para_text.strip():
                    p = doc.add_paragraph()
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_after = Pt(6)
                    run = p.add_run(para_text.strip())
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)

        if sec_bullets:
            for b in sec_bullets:
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_after = Pt(4)
                brun = bp.add_run(b)
                brun.font.name = "Calibri"
                brun.font.size = Pt(11)

        if sec_table and isinstance(sec_table, list) and len(sec_table) > 0:
            # Table formatting
            rows = len(sec_table)
            cols = len(sec_table[0])
            table = doc.add_table(rows=rows, cols=cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Table Grid'
            for r_idx, row in enumerate(sec_table):
                for c_idx, val in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = str(val)
                    p = cell.paragraphs[0]
                    if r_idx == 0:
                        p.runs[0].font.bold = True
                        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                        shading_elm = parse_xml(r'<w:shd {} w:fill="1B4F72"/>'.format(nsdecls('w')))
                        cell._tc.get_or_add_tcPr().append(shading_elm)

    doc.save(output_path)
    return output_path


def generate_pptx(title: str, subtitle: str, slides_data: List[Dict[str, Any]], output_path: str) -> str:
    """Generates a professional PowerPoint presentation."""
    if Presentation is None:
        raise ImportError("python-pptx is not installed.")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)  # 16:9 Widescreen
    prs.slide_height = PptxInches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Slide 1: Cover Slide
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(
        1, 0, 0, prs.slide_width, prs.slide_height  # 1 = MSO_SHAPE.RECTANGLE
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = PptxRGBColor(16, 44, 87)
    bg.line.color.rgb = PptxRGBColor(16, 44, 87)

    # Title box
    tx_box = slide.shapes.add_textbox(PptxInches(1.5), PptxInches(2.2), PptxInches(10.3), PptxInches(2.5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = PptxPt(40)
    p.font.bold = True
    p.font.color.rgb = PptxRGBColor(255, 255, 255)
    p.font.name = "Calibri"

    p2 = tf.add_paragraph()
    p2.text = subtitle or "PT Amsyong AI Solution Presentation"
    p2.font.size = PptxPt(20)
    p2.font.italic = True
    p2.font.color.rgb = PptxRGBColor(210, 225, 245)
    p2.font.name = "Calibri"

    # Subsequent Content Slides
    for item in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Header background banner
        header_bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, PptxInches(1.2))
        header_bar.fill.solid()
        header_bar.fill.fore_color.rgb = PptxRGBColor(24, 76, 120)
        header_bar.line.color.rgb = PptxRGBColor(24, 76, 120)

        # Slide Title
        t_box = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(0.2), PptxInches(11.5), PptxInches(0.8))
        t_frame = t_box.text_frame
        tp = t_frame.paragraphs[0]
        tp.text = item.get("title", "Slide Title")
        tp.font.size = PptxPt(28)
        tp.font.bold = True
        tp.font.color.rgb = PptxRGBColor(255, 255, 255)
        tp.font.name = "Calibri"

        # Content Box
        c_box = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(1.6), PptxInches(11.5), PptxInches(5.2))
        c_frame = c_box.text_frame
        c_frame.word_wrap = True

        main_text = item.get("content", "")
        if main_text:
            p_main = c_frame.paragraphs[0]
            p_main.text = main_text
            p_main.font.size = PptxPt(18)
            p_main.font.color.rgb = PptxRGBColor(40, 40, 40)
            p_main.font.name = "Calibri"

        bullets = item.get("bullets", [])
        for b_text in bullets:
            pb = c_frame.add_paragraph()
            pb.text = "• " + b_text
            pb.font.size = PptxPt(17)
            pb.font.color.rgb = PptxRGBColor(50, 50, 50)
            pb.font.name = "Calibri"

    prs.save(output_path)
    return output_path
